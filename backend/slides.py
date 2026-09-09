from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pathlib import Path
from typing import Any


def create_presentation(
    title: str,
    slides: list[dict],
    output_path: str,
    theme: str = "default",
) -> str:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title  # type: ignore[union-attr]
    ph = slide.placeholders[1]
    tf = ph.text_frame  # type: ignore[union-attr]
    if tf is not None:
        tf.text = ""

    for slide_data in slides:
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = slide_data.get("title", "")  # type: ignore[union-attr]

        body = slide.placeholders[1]
        tf = body.text_frame  # type: ignore[union-attr]
        tf.clear()

        for i, item in enumerate(slide_data.get("content", [])):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = item
            p.font.size = Pt(18)
            p.space_after = Pt(6)

    prs.save(output_path)
    return output_path


def presentation_from_markdown(markdown_text: str, output_path: str) -> str:
    lines = markdown_text.split("\n")
    title = "Apresentação"
    slides: list[dict] = []
    current_slide: dict | None = None

    for line in lines:
        line = line.strip()
        if not line:
            continue
        if line.startswith("# "):
            title = line[2:]
        elif line.startswith("## "):
            if current_slide:
                slides.append(current_slide)
            current_slide = {"title": line[2:], "content": []}
        elif line.startswith("- ") or line.startswith("* "):
            if current_slide is None:
                current_slide = {"title": "", "content": []}
            current_slide["content"].append(line[2:])

    if current_slide:
        slides.append(current_slide)

    return create_presentation(title, slides, output_path)