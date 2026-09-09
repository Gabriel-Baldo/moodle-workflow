from __future__ import annotations

import html as html_module
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    Document = None  # type: ignore[misc]
    Pt = None  # type: ignore[misc]


def markdown_to_html(md: str) -> str:
    lines = md.split("\n")
    out = []
    in_code = False
    for line in lines:
        if line.startswith("```"):
            in_code = not in_code
            if in_code:
                out.append("<pre><code>")
            else:
                out.append("</code></pre>")
            continue
        if in_code:
            out.append(line); continue
        if line.startswith("# "):
            out.append(f"<h1>{_eh(line[2:])}</h1>")
        elif line.startswith("## "):
            out.append(f"<h2>{_eh(line[2:])}</h2>")
        elif line.startswith("### "):
            out.append(f"<h3>{_eh(line[3:])}</h3>")
        elif line.startswith("- ") or line.startswith("* "):
            out.append(f"<li>{_eh(line[2:])}</li>")
        elif line.strip():
            out.append(f"<p>{_eh(line)}</p>")
    return "\n".join(out)


def _eh(text: str) -> str:
    return html_module.escape(text)


def format_pdf(content_markdown: str, output_path: str) -> str:
    try:
        from weasyprint import HTML  # type: ignore[import-untyped]
    except ImportError:
        html_path = output_path.rsplit(".", 1)[0] + ".html"
        Path(html_path).write_text(_html_template(content_markdown), encoding="utf-8")
        print(f"weasyprint não instalado. Salvo como HTML: {html_path}", flush=True)
        return html_path
    HTML(string=_html_template(content_markdown)).write_pdf(output_path)
    return output_path


def _html_template(body: str) -> str:
    return f"""<!DOCTYPE html><html><head><meta charset="utf-8">
<style>body{{font-family:'DejaVu Serif',serif;font-size:12pt;line-height:1.6;margin:2cm}}
h1{{font-size:18pt;border-bottom:2px solid #333;padding-bottom:4pt}}
h2{{font-size:14pt;margin-top:18pt}}
code{{background:#f4f4f4;padding:2pt 4pt}}
pre{{background:#f4f4f4;padding:8pt;border-radius:4pt}}</style></head><body>{body}</body></html>"""


def format_docx(content_markdown: str, output_path: str) -> str:
    if not DOCX_AVAILABLE or Document is None:
        print("Erro: python-docx não disponível. Instale: pip install python-docx", flush=True)
        raise SystemExit(1)
    doc = Document()  # type: ignore[attr-defined]
    style = doc.styles["Normal"]
    style.font.name = "Calibri"  # type: ignore[attr-defined]
    style.font.size = Pt(11)  # type: ignore[attr-defined]
    for line in content_markdown.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[2:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[3:], level=3)
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("```"):
            continue
        else:
            doc.add_paragraph(line)
    doc.save(output_path)
    return output_path